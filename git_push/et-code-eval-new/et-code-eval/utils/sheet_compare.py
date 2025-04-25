import os
import re
import threading
import time
import xlwings as xw
import openpyxl
from func_timeout import func_timeout, FunctionTimedOut

from utils.cells_compare import DeepSearchCells
from utils.keyword_ignore_map.mapping_dict import SKIP_SHEET_ATTR, SKIP_CELL_ATTR


class DeepSearch:
    def __init__(self, strict=False, skip_attr=SKIP_SHEET_ATTR):
        self.dsc = DeepSearchCells(skip_attr=SKIP_CELL_ATTR)
        self.SKIP_ATTR = skip_attr
        self.strict = strict

    def get_all_attributes(self, obj):
        attributes = {}
        stack = [(obj, '')]

        while stack:
            current_obj, path = stack.pop()

            ### 判断list, tuple, set, dict都要逐个取出来重新放到队列，并记录该属性轨迹 ###
            ### 判断如果是值，直接记录路径并添加到属性 ###
            ### 既不是值也不是list, tuple, set, dict，说明是对象，所以需要遍历属性attr_name ###
            if isinstance(current_obj, (list, tuple, set)):
                for index, item in enumerate(current_obj):
                    stack.append((item, f"{path}[{index}]"))
            elif isinstance(current_obj, dict):
                for key, value in current_obj.items():
                    stack.append((value, f"{path}[{key}]"))
            elif isinstance(current_obj, (int, float, str, bool)):
                attributes[path] = current_obj
            else:
                for attr_name in dir(current_obj):
                    if attr_name.startswith('__') and attr_name.endswith('__'):
                        continue
                    ### 自定义跳过的属性 ###
                    if attr_name in self.SKIP_ATTR:
                        continue

                    ### 避免无定义属性 ###
                    try:
                        attr_value = getattr(current_obj, attr_name)
                    except:
                        continue
                    ### 跳过方法功能函数 ###
                    if callable(attr_value): continue

                    ### 记录当前属性轨迹 ###
                    full_path = f"{path}.{attr_name}" if path else attr_name

                    # print(full_path)
                    ### 判断list, tuple, set, dict都要重新放到队列，并记录该属性轨迹 ###
                    ### 如果含有__dict__属性，说明是对象，同样需要继续放回队列，下次轮到后近一步深入遍历这个新的子对象的属性 ###
                    ### 直到最后没有这些属性，都成了值，说明搜索完成，记录到attributes字典中 ###
                    if isinstance(attr_value, (list, tuple, set, dict)):
                        stack.append((attr_value, full_path))
                    elif hasattr(attr_value, '__dict__'):
                        stack.append((attr_value, full_path))
                    else:
                        if full_path is not None:
                            attributes[full_path] = attr_value

        return attributes

    def replace_excel(self, file_path):
        print("文件格式转化 xls --->>> xlsx...")
        app = xw.App(visible=False, add_book=False)
        example = app.books.open(file_path)
        example.save(f"{file_path}x")
        example.close()
        app.quit()
        new_path = f"{file_path}x"
        os.remove(file_path)
        return new_path

    def load_workbook_with_timeout(self, file_path, timeout):

        try:
            ws = func_timeout(timeout, openpyxl.load_workbook, (file_path,))
            return ws
        except FunctionTimedOut:
            print(f'读取样张{file_path}超时')
            raise TimeoutError

    def filter_elements_within_range(self, data, max_row, max_col):
        filtered_data = {k: v for k, v in data.items() if k[0] <= max_row and k[1] <= max_col}
        return filtered_data

    def excel_to_attribution_dict(self, excel_file, sheet_name=None):

        if excel_file.endswith('.xlsx'):
            wb = self.load_workbook_with_timeout(excel_file, 5)
        elif excel_file.endswith('.xls'):
            # try:
            new_path = self.replace_excel(excel_file)
            wb = self.load_workbook_with_timeout(new_path, 5)
            # except:
            #     return {}, []

        ws = wb.active

        if sheet_name:
            ws = wb[sheet_name]

        # if not self.strict:
        #     if ws.max_row>1000 or ws.max_column>500:
        #         raise TypeError

        sheet_name = ws.title

        workbook_sheetnames = wb.sheetnames

        all_attributes1 = self.get_all_attributes(ws)
        all_attributes1["workbook_sheetnames"]=workbook_sheetnames
        res = ws._cells
        res = self.filter_elements_within_range(res, 1000, 500)
        return all_attributes1, res, sheet_name

    def compare_dict(self, all_attributes1, all_attributes2, code):
        differences = []
        # all_attributes1_all = self.merge_nested_dict(all_attributes1)
        # all_attributes2_all = self.merge_nested_dict(all_attributes2)

        for elem in set(all_attributes1.keys()).union(all_attributes2.keys()):
            if elem in all_attributes1.keys() and elem in all_attributes2.keys():
                # 跳过比较不对行高列宽操作的case对属性的比较
                # 如果 code 中没有 "height"、"width" 或 "AutoFit"，那么进入下一步。
                # 如果 elem 中同时包含 ["column_dimensions", "row_dimensions"] 中的任意一个和 ["height", "width", "ht"] 中的任意一个，那么执行 continue，跳过当前循环的剩余部分，进入下一次循环。
                if any([True for i in ["column_dimensions", "row_dimensions"] if i in elem]) and any(
                        [True for i in [".height", ".width", ".ht"] if i in elem]):
                    if not any([True for i in ["height", "width", "AutoFit"] if i in code]):
                            continue
                    else:
                        sheet1_elem = all_attributes1[elem]
                        sheet2_elem = all_attributes2[elem]
                        try:
                            if sheet1_elem==sheet2_elem or abs(sheet1_elem-sheet2_elem)<0.1:
                                continue
                        except:
                            if ".h" in elem:
                                sheet1_elem = sheet1_elem if sheet1_elem else all_attributes1["sheet_format.defaultRowHeight"]
                                sheet2_elem = sheet2_elem if sheet2_elem else all_attributes2["sheet_format.defaultRowHeight"]
                            elif ".w" in elem:
                                sheet1_elem = sheet1_elem if sheet1_elem else all_attributes1["sheet_format.defaultColWidth"]
                                sheet2_elem = sheet2_elem if sheet2_elem else all_attributes2["sheet_format.defaultColWidth"]
                            differences.append(f"*{elem}*: ({sheet1_elem}) -> ({sheet2_elem})")
                            continue


                sheet1_elem = all_attributes1[elem]
                sheet2_elem = all_attributes2[elem]

                if sheet1_elem==sheet2_elem:
                    continue
                else:
                    # if "row_dimensions" in elem and sheet1_elem==17.25 and  sheet2_elem==20.25:
                    #     continue
                    differences.append(f"*{elem}*: ({sheet1_elem}) -> ({sheet2_elem})")
            else:
                ### 如果要求不严格，一个没有属性，暂时跳过；如果要求严格，这种也要加上###
                if not self.strict:
                    continue
                else:
                    try:
                        differences.append(f"*{elem}*: ({all_attributes1[elem]}) -> (Missing Attribution)")
                    except:
                        differences.append(f"*{elem}*: (Missing Attribution) -> ({all_attributes2[elem]})")

        return differences

    def run_compare(self, excel_file1, excel_file2, code_dict):
        try:
            all_attributes1, cellsList1, sheet_name = self.excel_to_attribution_dict(excel_file1)
            all_attributes2, cellsList2, _ = self.excel_to_attribution_dict(excel_file2, sheet_name)
        except TypeError as typeerror:
            return ([f"校验样张解析错误 (超长 | openpyxl读取失败) :{typeerror}"], [])
        except KeyError as keyerror:
            return ([f"校验样张解析错误:{keyerror}"], [])
        except TimeoutError as timeerror:
            return ([f"样张读取超时:{timeerror}"], [])

        cur_index = os.path.basename(os.path.dirname(excel_file1))

        diff_ws = self.compare_dict(all_attributes1, all_attributes2, code_dict[cur_index])

        diff_cells = self.dsc.run_compare_cells(cellsList1, cellsList2, cur_index)

        return (diff_ws, diff_cells)


if __name__ == "__main__":
    ds = DeepSearch()

    file1 = r"C:\Users\Bolt\PycharmProjects\et-code-eval\case\test0.1k\run_code_results\06cb3ff8235f5aa2b81662c1900eaf06\06cb3ff8235f5aa2b81662c1900eaf06_原始样张.xlsx"
    file2 = r"C:\Users\Bolt\PycharmProjects\et-code-eval\case\test0.1k\run_code_results\06cb3ff8235f5aa2b81662c1900eaf06\06cb3ff8235f5aa2b81662c1900eaf06_执行样张.xlsx"

    res = ds.run_compare(file1, file2)


    print("表格属性差异：")
    for item in res[0]:
        if item:   print("\t", item)
    print("单元格属性差异：")
    for item in res[1]:
        if item:   print("\t", item)
