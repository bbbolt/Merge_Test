# -*- coding: utf-8 -*-
# @Time : 2024/6/24 下午3:55
# @Author : sunyuzhao
# @Email : sunyuzhao@wps.cn
# @File : rag.py
import openpyxl
from collections import defaultdict

def loadFromFile():

    workbook = openpyxl.load_workbook('core/case.xlsx')
    ex = defaultdict(dict)
    # 遍历每个Sheet
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        # 遍历每行数据
        key = ''
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if (row[0] is None or row[1] is None or row[2] is None) or row[4] is None:
                continue

            ex[row[0]]['table_structure'] = row[1]
            ex[row[1]]['answer'] = row[2]

    return ex