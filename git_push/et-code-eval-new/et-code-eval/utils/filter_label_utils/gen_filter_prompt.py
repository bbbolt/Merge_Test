from transformers import AutoModelForCausalLM, AutoTokenizer, TextStreamer
from tqdm import tqdm
import json
import torch
import re
import argparse
import os
import openpyxl

def distill_sample_item(data_path):

    wb = openpyxl.load_workbook(data_path)
    ws = wb.active

    # 获取列名所在的行
    header = [cell.value for cell in ws[1]]

    # 找到列名对应的列索引

    md5_index = header.index('表格名(md5)')
    prompt_index = header.index('表格描述')
    code_index = header.index('代码')
    diff_index = header.index('差异信息')
    flag_index = header.index('执行前后是否有变化')

    data_list = []
    # 遍历每一行（从第二行开始，因为第一行是列名）
    for row in ws.iter_rows(min_row=2, values_only=True):
        md5 = row[md5_index]
        prompt = row[prompt_index]
        code = row[code_index]
        diff = row[diff_index]
        flag = row[flag_index]
        if not code: continue
        if not code.strip(): continue
        if not diff: continue
        if not diff.strip(): continue
        # 创建字典并添加到列表中
        data_dict = {
            "input": prompt,
            "target": code,
            "md5": md5 if type(md5)==str else str(int(md5)),
            "diff": diff,
            "flag": flag
        }
        data_list.append(data_dict)
    

    res = []
    for json_dict in tqdm(data_list, desc="组合表格描述+用户需求+属性差异 ---> Prompts："):
        input = f"""现在我使用WPS-JS宏代码对表格进行了处理，但不确定执行是否正确，你需要根据我给定的代码执行前后openpyxl解析的属性差异，判断JS宏代码是否满足用户需求。

表格描述及对应用户需求：
{json_dict["input"]}

WPS-JS宏代码：
{json_dict["target"].split("//$$$")[0]}

代码执行前后openpyxl属性差异格式为--属性: (原样张属性值 --->>> 执行后样张属性值)：
{json_dict["diff"]}

输出内容结构:
{{
"FLAG": True or False,
"ANALYSIS": 添加原因
}}

注意：请严格按照上述格式填写，否则将无法通过测试。
"""
        res.append({
            "input": input,
            "md5": json_dict["md5"],
            "flag": json_dict["flag"]
        })

    return res

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='gen sheet')
    parser.add_argument('--data_path', type=str, help='data_path', default="/home/kas/x_gen_code_sheet/filter_label_code/input_xlsx/测评_241225_code.xlsx")
    args = parser.parse_args()
    distill_sample_item(args.data_path)

